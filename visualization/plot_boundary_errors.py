import os
import subprocess
import sys

def install_and_import(package):
    try:
        import docx
    except ImportError:
        print(f"Installing {package}...")
        subprocess.check_call([sys.executable, "-m", "pip", "install", package])
    finally:
        globals()["docx"] = __import__("docx")

install_and_import("python-docx")

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_report(output_path):
    doc = Document()
    
    # Title
    title = doc.add_heading('Analysis of High Boundary Errors in Physics-Informed Neural Networks (PINNs) for Groundwater Modeling', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Introduction
    doc.add_heading('1. Introduction', level=1)
    p = doc.add_paragraph(
        "The observation of high errors concentrated at the boundaries of the study area is a well-documented phenomenon when applying Physics-Informed Neural Networks (PINNs) to regional groundwater flow modeling. While traditional numerical solvers like MODFLOW strictly enforce boundary conditions (BCs), PINNs approximate the solution using continuous, differentiable functions. This discrepancy in how boundary conditions are mathematically treated often leads to reduced accuracy at the model’s periphery. The issue stems fundamentally from the optimization dynamics and formulation used during the PINN's training phase."
    )
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Section 2
    doc.add_heading('2. Soft vs. Hard Boundary Constraints and Loss Imbalance', level=1)
    p = doc.add_paragraph(
        "The most significant factor contributing to boundary errors is the nature of the loss function. In standard numerical models, Dirichlet (specified head) or Neumann (specified flux) boundaries are treated as \"hard constraints\"—the solver is mathematically forced to satisfy them exactly. In contrast, standard PINNs implement boundary conditions as \"soft constraints.\" The boundary loss is merely one penalty term added to the partial differential equation (PDE) residual loss and the empirical data loss.\n\n"
        "During training, the optimizer seeks to minimize the global loss landscape. If the weights assigned to these competing loss components are not perfectly balanced, the network may effectively \"ignore\" the boundary conditions to achieve a lower overall PDE residual in the interior domain. This gradient pathology, where the gradients of the PDE loss dominate the gradients of the boundary loss, results in the PINN failing to learn the correct boundary states."
    )
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Section 3
    doc.add_heading('3. Geometric Complexity and Collocation Point Density', level=1)
    p = doc.add_paragraph(
        "The boundary of the modeled basin is highly irregular, derived from a jagged rasterized active cell configuration (where ibound > 0). Neural networks tend to struggle with complex, non-convex spatial domains unless guided by extremely dense sampling. If the collocation points used to evaluate the boundary conditions are too sparse or uniformly sampled across the overall domain, the network will lack the localized spatial resolution required to learn the complex geometric edges. Consequently, the PINN smooths over these boundary intricacies, causing high localized errors."
    )
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Section 4
    doc.add_heading('4. Spectral Bias and Sharp Gradients', level=1)
    p = doc.add_paragraph(
        "Neural networks exhibit a mathematical property known as \"spectral bias,\" meaning they preferentially learn low-frequency, smooth, continuous functions before capturing high-frequency, complex features. Model boundaries typically represent abrupt transitions, such as no-flow bedrock interfaces or varying river stages, which manifest as sharp hydrogeological gradients. Standard multi-layer perceptron (MLP) architectures struggle to resolve these sharp gradients at the edges because they default to a smoother, generalized surface that fits the interior well but deviates significantly at the rigid boundaries."
    )
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Section 5
    doc.add_heading('5. Strategies for Remediation in PINN Training', level=1)
    p = doc.add_paragraph(
        "To mitigate these high boundary errors in subsequent training iterations, several modifications to the PINN methodology can be employed:"
    )
    
    # Bullet points
    doc.add_paragraph("Adaptive Loss Weighting: Implement dynamic weight adjustment algorithms (e.g., Learning Rate Annealing or ReLoBRaLo) that automatically scale up the penalty weight of the boundary loss if the boundary error remains disproportionately high compared to the interior PDE loss.", style='List Bullet')
    doc.add_paragraph("Spatial Oversampling: Significantly increase the density of collocation points exclusively along the external boundary edges and near complex topological features. This forces the optimizer to pay more attention to the boundary.", style='List Bullet')
    doc.add_paragraph("Hard Constraint Formulations: Modify the neural network architecture using an exact boundary enforcement technique (e.g., using distance functions). This involves multiplying the network's output by an exact spatial distance function that evaluates to zero at the boundary, ensuring boundary conditions are mathematically guaranteed rather than merely penalized.", style='List Bullet')

    # Conclusion
    doc.add_heading('6. Conclusion', level=1)
    p = doc.add_paragraph(
        "The elevated error at the boundary is not an anomaly but a direct consequence of the soft-constraint optimization paradigm of standard PINNs. By adjusting the loss weighting strategy, refining boundary point sampling, or transitioning to hard-constraint architectures, the boundary accuracy can be systematically brought into alignment with the high performance observed in the interior model domain."
    )
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Save Document
    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")

if __name__ == "__main__":
    workspace = r'E:\mayank\varuna_refinement\try_1_master'
    output_dir = os.path.join(workspace, 'drop_1000', 'reports')
    os.makedirs(output_dir, exist_ok=True)
    
    output_file = os.path.join(output_dir, 'PINN_Boundary_Error_Analysis.docx')
    create_report(output_file)
